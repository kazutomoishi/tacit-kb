#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
make_blind_eval.py  —  盲検・対応ありの評価シートを生成
  前提: extract_batch_control.py を実行済み（cards に source='llm_control' がある）
  使い方:
      python make_blind_eval.py                    # 各条件40枚・計80枚
      python make_blind_eval.py 30                 # 各条件30枚・計60枚

  設計:
    - 対応あり : 両条件のカードを持つ論文から、1論文につき各条件1枚ずつ
    - 層別     : 理論あり側カードの重要度で 高:中:低 = 16:14:10（40枚時）
    - 盲検     : 両条件を混ぜてシャッフルし、通し番号を振り直す
                 条件・card_id は評価者用ブックに一切残さない

  出力:
      tacit_card_eval_blind.xlsx   ← 評価者に渡す（条件が分からない）
      blind_key.csv                ← 対応表。評価者には渡さないこと
"""
import sqlite3, sys, random, csv, re
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation

DB = "tacit_kb.db"
N_PER_ARM = int(sys.argv[1]) if len(sys.argv) > 1 else 40
SEED = 20260806

HDR_FILL = PatternFill("solid", start_color="FF1F4E79")
EX_FILL = PatternFill("solid", start_color="FFEDEDED")
IN_FILL = PatternFill("solid", start_color="FFFFF3A3")
THIN = Side(style="thin", color="FFBFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

COLS = [
    ("no", 5), ("card_id", 8), ("重要度\n(モデル)", 9), ("出典文 source_passage", 46),
    ("term", 20), ("why_vague\nなぜ曖昧か", 34), ("body_sense\n身体感覚", 30),
    ("concrete_steps\n明示的手順", 40), ("failure_mode\n失敗様態", 34),
    ("①忠実性\n1-4", 7), ("②捏造\n有/無", 7), ("③妥当性\n1-4", 7),
    ("④有用性\n1-4", 7), ("⑤重要度ラベル\n適切/過大/過小", 12), ("コメント", 30),
]

DESC = [
    ("暗黙知カード 忠実性評価シート  —  評価者 A", True),
    ("", False),
    ("目的：本手法（LLMが生成した「暗黙知カード」）の出力品質を、2名の評価者が独立に評価します。", False),
    ("評価者Aと評価者Bは相談せず、各自でこの「評価」タブを記入してください（独立性が一致度κの前提です）。", False),
    (f"対象：カード{N_PER_ARM*2}枚。重要度でモデルが付けたラベルに基づき層別抽出。所要の目安は1枚2〜3分です。", False),
    ("", False),
    ("記入する列（黄色）：①忠実性 ②捏造 ③妥当性 ④有用性 ⑤重要度ラベル評価 と コメント。カード内容の列（白）は編集しないでください。", False),
    ("", False),
    ("① 忠実性（1-4）：出典文（source_passage）に照らし、カードが出典にない値・手順・条件を作り出していないか。", False),
    ("　4=完全に出典に忠実（捏造なし）／3=概ね忠実（軽微な一般的補足のみ）／2=やや不十分（根拠不明の具体値・手順あり）／1=不十分（明らかな捏造・出典と矛盾）", False),
    ("② 捏造（有/無）：出典に無い具体的な数値・手順・条件が含まれるなら「有」。①と別に明示してください。", False),
    ("③ 理論的妥当性（1-4）：why_vague・body_sense の説明（二重過程／スキーマ／SECI の読み）が妥当か。4=妥当〜1=不当。", False),
    ("④ 有用性（1-4）：concrete_steps・failure_mode が、実務者が判断を再現するのに役立つか。4=有用〜1=無用。", False),
    ("⑤ 重要度ラベル評価：モデルが付けた重要度（C列）が妥当か。 適切／過大（重要すぎ）／過小（軽視しすぎ）。", False),
    ("", False),
    ("出典文が「(本文一致なし…)」の行：語句が本文と完全一致せず自動抽出できなかったカードです。PMCIDから原文を参照するか、判断が難しければ空欄のままで構いません。", False),
    ("", False),
    ("集計：幻覚率＝「①忠実性≤2 または ②捏造=有」の割合。評価者間一致は忠実性の可否（≥3か否か）でCohenのκを算出。", False),
    ("注意：本サンプルは重要度で均等寄りに抽出しているため、コーパス全体の率を出す場合は重要度比で再重み付けしてください。", False),
    ("", False),
    ("本シートのカードは複数の生成設定から無作為に混合し、順序をシャッフルしています。設定の別は評価終了まで開示されません。", False),
    ("設定ごとに評価が偏らないよう、各カードは単独で、出典文のみに照らして評価してください。", False),
]


def source_passage(c, pid, term):
    """term を同一論文の snippet と照合。無ければ PMCID+タイトルのフォールバック。"""
    t = (term or "").strip()
    if t:
        rows = c.execute(
            "SELECT snippet FROM candidate_terms WHERE paper_id=? AND snippet LIKE ?",
            (pid, f"%{t}%")).fetchall()
        if rows:
            return " ".join(rows[0][0].split())
    r = c.execute("SELECT pmcid, title FROM papers WHERE id=?", (pid,)).fetchone()
    pmcid, title = (r or ("", ""))
    return f"(本文一致なし／原文を参照) PMCID {pmcid}: {title}"


def main():
    random.seed(SEED)
    conn = sqlite3.connect(DB)
    c = conn.cursor()

    srcs = {r[0] for r in c.execute("SELECT DISTINCT source FROM cards")}
    if "llm_control" not in srcs:
        sys.exit(f"対照条件のカードがありません。先に extract_batch_control.py を実行してください。\n  現在の source: {sorted(srcs)}")

    # 両条件を持つ論文
    a = {r[0] for r in c.execute("SELECT DISTINCT paper_id FROM cards WHERE source='llm'")}
    b = {r[0] for r in c.execute("SELECT DISTINCT paper_id FROM cards WHERE source='llm_control'")}
    common = sorted(a & b)
    if len(common) < N_PER_ARM:
        sys.exit(f"両条件が揃う論文が {len(common)} 本しかありません（必要 {N_PER_ARM} 本）。")

    # 理論あり側カードを重要度で層別 → 論文を選ぶ
    ratio = {"高重要度": 0.40, "中重要度": 0.35, "低重要度": 0.25}
    quota = {k: round(N_PER_ARM * v) for k, v in ratio.items()}
    quota["高重要度"] += N_PER_ARM - sum(quota.values())

    by_crit = {k: [] for k in ratio}
    ph = ",".join("?" * len(common))
    for cid, pid, crit in c.execute(
            f"SELECT id, paper_id, criticality FROM cards WHERE source='llm' AND paper_id IN ({ph})", common):
        if crit in by_crit:
            by_crit[crit].append((cid, pid))
    for v in by_crit.values():
        random.shuffle(v)

    chosen, used_papers = [], set()
    for crit, need in quota.items():
        for cid, pid in by_crit[crit]:
            if need <= 0:
                break
            if pid in used_papers:
                continue
            used_papers.add(pid)
            chosen.append((cid, pid, crit))
            need -= 1
        if need > 0:
            print(f"  警告: {crit} が {need} 枚不足（他層で補填します）")

    # 不足分を任意の層から補填
    if len(chosen) < N_PER_ARM:
        pool = [(cid, pid, cr) for cr, v in by_crit.items() for cid, pid in v if pid not in used_papers]
        random.shuffle(pool)
        for cid, pid, cr in pool:
            if len(chosen) >= N_PER_ARM:
                break
            used_papers.add(pid)
            chosen.append((cid, pid, cr))

    # 各論文から対照条件を1枚
    rows = []
    for cid, pid, crit in chosen:
        ca = c.execute("""SELECT id,criticality,term,why_vague,body_sense,concrete_steps,failure_mode
                          FROM cards WHERE id=?""", (cid,)).fetchone()
        cbs = c.execute("""SELECT id,criticality,term,why_vague,body_sense,concrete_steps,failure_mode
                           FROM cards WHERE source='llm_control' AND paper_id=?""", (pid,)).fetchall()
        if not cbs:
            continue
        cb = random.choice(cbs)
        rows.append(("theory", pid, ca))
        rows.append(("control", pid, cb))

    random.shuffle(rows)
    print(f"対象論文 {len(chosen)} 本 / カード {len(rows)} 枚（理論あり {sum(1 for r in rows if r[0]=='theory')} ・"
          f"理論なし {sum(1 for r in rows if r[0]=='control')}）")

    # ---------------- ブック作成 ----------------
    wb = Workbook()
    ws0 = wb.active
    ws0.title = "説明"
    ws0.column_dimensions["A"].width = 110
    for i, (txt, bold) in enumerate(DESC, 1):
        cell = ws0.cell(row=i, column=1, value=txt)
        cell.font = Font(name="Arial", size=11, bold=bold)
        cell.alignment = Alignment(wrap_text=True, vertical="top")

    ws = wb.create_sheet("評価")
    for j, (name, w) in enumerate(COLS, 1):
        cell = ws.cell(row=1, column=j, value=name)
        cell.font = Font(name="Arial", size=10, bold=True, color="FFFFFFFF")
        cell.fill = HDR_FILL
        cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        cell.border = BORDER
        ws.column_dimensions[chr(64 + j)].width = w
    ws.row_dimensions[1].height = 30

    ex = ["例", "(例)", "高",
          "Incubate on ice for exactly the time needed…", "on ice",
          "「on ice」だけで温度・時間の判断が省略されている", "容器の冷たさが持続する感触",
          "4°Cで保持し使用直前まで氷上に置く", "室温放置で酵素活性が落ちる",
          4, "無", 4, 3, "適切", "例：記入前にこの行を削除"]
    for j, v in enumerate(ex, 1):
        cell = ws.cell(row=2, column=j, value=v)
        cell.font = Font(name="Arial", size=9)
        cell.fill = EX_FILL
        cell.alignment = Alignment(wrap_text=(j in (4, 6, 7, 8, 9, 15)), vertical="top")
        cell.border = BORDER
    ws.row_dimensions[2].height = 24

    key = []
    for i, (arm, pid, card) in enumerate(rows, 1):
        cid, crit, term, why, body, steps, fail = card
        r = i + 2
        vals = [i, "", (crit or "").replace("重要度", ""),
                source_passage(c, pid, term), term, why, body, steps, fail,
                None, None, None, None, None, None]
        for j, v in enumerate(vals, 1):
            cell = ws.cell(row=r, column=j, value=v)
            cell.font = Font(name="Arial", size=9)
            cell.alignment = Alignment(wrap_text=(j in (4, 5, 6, 7, 8, 9, 15)), vertical="top")
            cell.border = BORDER
            if j >= 10:
                cell.fill = IN_FILL
        ws.row_dimensions[r].height = 70
        key.append({"no": i, "arm": arm, "card_id": cid, "paper_id": pid, "criticality": crit})

    last = len(rows) + 2
    for formula, rng in [('"1,2,3,4"', f"J3:J{last} L3:L{last} M3:M{last}"),
                         ('"有,無"', f"K3:K{last}"),
                         ('"適切,過大,過小"', f"N3:N{last}")]:
        dv = DataValidation(type="list", formula1=formula, allow_blank=True)
        ws.add_data_validation(dv)
        dv.sqref = rng
    ws.freeze_panes = "J2"

    wb.save("tacit_card_eval_blind.xlsx")

    with open("blind_key.csv", "w", newline="", encoding="utf-8-sig") as f:
        w = csv.DictWriter(f, fieldnames=["no", "arm", "card_id", "paper_id", "criticality"])
        w.writeheader()
        w.writerows(key)

    conn.close()
    print("\n出力:")
    print("  tacit_card_eval_blind.xlsx  ← 評価者に渡す")
    print("  blind_key.csv               ← 対応表（評価者には渡さないこと）")


if __name__ == "__main__":
    main()
